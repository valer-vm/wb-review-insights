"""Тесты report.py на сохранённых JSON — без субагентов и сети."""
import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

import report

RESULTS_DIR = Path(__file__).parent / "fixtures" / "results"
NM, OTHER = 368686013, 111222333


def cells(ws):
    return [[c for c in row] for row in ws.iter_rows(values_only=True)]


def find_row(ws, first):
    return next(r for r in ws.iter_rows(values_only=True) if r[0] == first)


@pytest.fixture
def run(tmp_path):
    (tmp_path / "data").mkdir()
    shutil.copytree(RESULTS_DIR, tmp_path / "results")
    for art, brand in ((NM, "BODO"), (OTHER, "Конкурент")):
        meta = {"article": art, "brand": brand, "name": "Шапка", "rating": "4.9", "total_feedbacks": 1938,
                "months": 12, "url": f"https://www.wildberries.ru/catalog/{art}/detail.aspx"}
        (tmp_path / "data" / f"{art}.meta.json").write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
    status = [
        {"input": str(NM), "article": NM, "status": "ok"},
        {"input": str(OTHER), "article": OTHER, "status": "ok"},
        {"input": "https://ozon.ru/1", "status": "error", "message": "не похоже на ссылку на товар WB"},
        {"input": "555", "article": 555, "status": "error", "message": "товар 555 не найден на WB"},
    ]
    (tmp_path / "data" / "_status.json").write_text(json.dumps(status, ensure_ascii=False), encoding="utf-8")
    return tmp_path


def build(run, log=None):
    out = run / "report.xlsx"
    info = report.build_report(run, out, log=log or (lambda s: None))
    return load_workbook(out), info


def test_sheets(run):
    wb, _ = build(run)
    assert wb.sheetnames == ["Свод", "Рекомендации", f"BODO ({NM})", f"Конкурент ({OTHER})", "Сбор"]


def test_summary_sheet_counts_and_examples(run):
    wb, _ = build(run)
    ws = wb["Свод"]
    row = find_row(ws, "Маломерит")
    assert row[2] == 3 + 6                       # упоминания сложены по конкурентам
    assert row[3] == "2 из 2"
    assert row[4] == f"«Очень тугая, уши больно, хотя по размеру» — BODO ({NM}), 4★"
    assert ws.cell(row=[r[0] for r in cells(ws)].index("Маломерит") + 1, column=1).fill.fgColor.rgb.endswith("FDE9E7")
    single = find_row(ws, "Грубая ткань")
    assert single[3] == "1 из 2" and single[5] is None
    assert find_row(ws, "Тёплая и плотная")[2] == 4 + 10
    assert "Конкурентов: 2" in ws["A2"].value and "32 негативных" in ws["A2"].value


def test_summary_sorted_by_spread_then_mentions(run):
    rows = [r[0] for r in cells(build(run)[0]["Свод"])]
    complaints = rows[rows.index("Тема") + 1: rows.index("Похвала") - 1]
    assert complaints == ["Маломерит", "Теряет цвет после стирки", "Грубая ткань"]


def test_recommendations_sheet(run):
    ws = build(run)[0]["Рекомендации"]
    flat = [v for r in cells(ws) for v in r if v]
    assert "Что улучшить в товаре" in flat and "Что подчеркнуть в карточке" in flat
    assert find_row(ws, "Проверить лекала по размерам")[2] == "Маломерит"
    assert any(isinstance(v, str) and v.startswith("Главные жалобы ниши") for v in flat)


def test_competitor_sheet(run):
    ws = build(run)[0][f"BODO ({NM})"]
    assert ws["A2"].hyperlink.target.endswith(f"/{NM}/detail.aspx")
    assert "проанализировано: 12 негативных, 12 позитивных" in ws["A3"].value
    row = find_row(ws, "Выгорает и теряет цвет после стирок")
    assert row[2] == 5 and row[3] == round(100 * 5 / 12)
    assert row[4].endswith("(2★)") and row[6] is not None


def test_status_sheet_lists_failures(run):
    ws = build(run)[0]["Сбор"]
    rows = cells(ws)[3:]
    assert rows[0][:3] == [NM, "собрано", "есть"]
    assert rows[2][1] == "ошибка" and rows[2][3].startswith("https://ozon.ru/1:")
    assert rows[3][:2] == [555, "ошибка"]


def test_invalid_result_is_skipped(run):
    (run / "results" / f"{OTHER}.json").write_text("{битый", encoding="utf-8")
    logs = []
    wb, info = build(run, logs.append)
    assert list(info["results"]) == [NM]
    assert any(f"{OTHER}.json" in m for m in logs)
    # summary ссылается на пропавшего конкурента — тоже пропущен, но отчёт строится
    assert info["summary"] is None and any("summary.json" in m for m in logs)
    assert "Сводки нет" in wb["Свод"]["A4"].value
    assert find_row(wb["Сбор"], OTHER)[2] == "нет (ошибка анализа)"


def test_missing_summary(run):
    (run / "results" / "summary.json").unlink()
    wb, info = build(run)
    assert info["summary"] is None and f"BODO ({NM})" in wb.sheetnames


def test_no_results_fails(run):
    for p in (run / "results").glob("*.json"):
        p.unlink()
    with pytest.raises(SystemExit, match="Нет ни одного"):
        build(run)


def test_sheet_title():
    used = set()
    assert report.sheet_title("A/B: [x]?", used) == "A B   x"
    long = "Очень длинное название бренда без пробелов 123"
    first, second = report.sheet_title(long, used), report.sheet_title(long, used)
    assert len(first) == 31 and len(second) <= 31 and first != second and second.endswith(" 2")


def test_main_and_console_summary(run, capsys):
    assert report.main(["--run", str(run)]) == 0
    out = capsys.readouterr().out
    assert (run / "report.xlsx").exists()
    assert "Маломерит: 9 / 2 — общая для ниши" in out
    assert "555: товар 555 не найден" in out and "https://ozon.ru/1" in out

    for p in (run / "results").glob("*.json"):
        p.unlink()
    assert report.main(["--run", str(run)]) == 1
