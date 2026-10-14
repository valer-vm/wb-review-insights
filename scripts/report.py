# /// script
# requires-python = ">=3.11"
# dependencies = ["pandas>=2.1", "openpyxl>=3.1"]
# ///
"""Excel-отчёт из JSON, которые написали субагенты.

  uv run scripts/report.py --run wb-insights/2026-10-15_1200

В папке запуска ожидаются data/ (meta.json и _status.json от collect.py)
и results/ (<артикул>.json от review-analyst, summary.json от insights-writer).
Листы: «Свод», «Рекомендации», по листу на конкурента, «Сбор».
Невалидные файлы пропускаются с предупреждением — отчёт строится из того, что есть.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

import schema

HEADER_FILL = PatternFill("solid", fgColor="E8E0F5")
NICHE_FILL = PatternFill("solid", fgColor="FDE9E7")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
WRAP = Alignment(wrap_text=True, vertical="top")
STATUS_RU = {"ok": "собрано", "error": "ошибка", "duplicate": "пропущено (дубль)", "empty": "нет отзывов"}


def load_run(run: Path, log=print) -> tuple[dict[int, dict], dict[int, dict], dict | None, list[dict]]:
    """results (только валидные), meta, summary (или None), status."""
    data, res_dir = run / "data", run / "results"
    status = _read_json(data / "_status.json") or []
    meta = {}
    for p in sorted(data.glob("*.meta.json")):
        m = _read_json(p)
        if m:
            meta[int(m["article"])] = m
    results = {}
    for p in sorted(res_dir.glob("*.json")):
        if p.name == "summary.json":
            continue
        r = _read_json(p)
        errs = ["невалидный JSON"] if r is None else schema.validate_result(r)
        if errs:
            log(f"Пропускаю {p.name}: {errs[0]}" + (f" (и ещё {len(errs) - 1})" if len(errs) > 1 else ""))
            continue
        results[int(r["article"])] = r
    summary = _read_json(res_dir / "summary.json")
    if summary is not None:
        errs = schema.validate_summary(summary, results)
        if errs:
            log(f"Пропускаю summary.json: {errs[0]}" + (f" (и ещё {len(errs) - 1})" if len(errs) > 1 else ""))
            summary = None
    return results, meta, summary, status


def _read_json(p: Path):
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return None


# ---------------------------------------------------------------- оформление


def _header(ws: Worksheet, row: int, titles: list[str]) -> None:
    for col, t in enumerate(titles, 1):
        c = ws.cell(row=row, column=col, value=t)
        c.font, c.fill, c.alignment = BOLD, HEADER_FILL, WRAP


def _row(ws: Worksheet, row: int, values: list, fill: PatternFill | None = None) -> None:
    for col, v in enumerate(values, 1):
        c = ws.cell(row=row, column=col, value=v)
        c.alignment = WRAP
        if fill:
            c.fill = fill


def _widths(ws: Worksheet, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _label(article: int, results: dict, meta: dict) -> str:
    r = results.get(article) or meta.get(article) or {}
    brand = (r.get("brand") or "").strip()
    return f"{brand} ({article})" if brand else str(article)


def sheet_title(text: str, used: set[str]) -> str:
    """Имя листа Excel: без []:*?/\\, до 31 символа, уникальное."""
    base = re.sub(r"[\[\]:*?/\\]", " ", text).strip()[:31] or "Лист"
    name, n = base, 2
    while name.lower() in used:
        suffix = f" {n}"
        name, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(name.lower())
    return name


# ---------------------------------------------------------------- листы


def write_summary_sheet(ws: Worksheet, results: dict, meta: dict, summary: dict | None) -> None:
    ws.title = "Свод"
    ws["A1"] = "Отзывы конкурентов: на что жалуются и за что хвалят"
    ws["A1"].font = TITLE
    total_neg = sum(r["reviews_analyzed"]["negative"] for r in results.values())
    total_pos = sum(r["reviews_analyzed"]["positive"] for r in results.values())
    ws["A2"] = (f"Конкурентов: {len(results)}. Проанализировано отзывов: {total_neg} негативных, "
                f"{total_pos} позитивных. Розовым выделены проблемы, общие для ниши.")
    _widths(ws, [34, 50, 13, 13, 50, 50, 50])
    row = 4
    if summary is None:
        ws.cell(row=row, column=1, value="Сводки нет (insights-writer не отработал) — смотрите листы конкурентов.").font = BOLD
        return
    for section, title in (("complaints", "Жалобы"), ("praise", "Похвала")):
        ws.cell(row=row, column=1, value=title).font = TITLE
        row += 1
        _header(ws, row, ["Тема", "Описание", "Упоминаний", "У скольких конкурентов", "Пример 1", "Пример 2", "Пример 3"])
        row += 1
        groups = []
        for g in summary[section]:
            st = schema.summary_stats(g, results, section)
            groups.append((st, g))
        groups.sort(key=lambda x: (-x[0]["competitors"], -x[0]["mentions"]))
        for st, g in groups:
            examples = [f"«{e['quote']}» — {_label(e['article'], results, meta)}"
                        + (f", {e['rating']}★" if e.get("rating") else "") for e in g["examples"]]
            fill = NICHE_FILL if section == "complaints" and st["niche_wide"] else None
            _row(ws, row, [g["theme"], g["description"], st["mentions"],
                           f"{st['competitors']} из {len(results)}", *examples], fill)
            row += 1
        row += 1


def write_recommendations_sheet(ws: Worksheet, summary: dict | None) -> None:
    ws["A1"] = "Что сделать продавцу"
    ws["A1"].font = TITLE
    _widths(ws, [40, 80, 40])
    if summary is None:
        ws["A3"] = "Рекомендаций нет (insights-writer не отработал)."
        return
    row = 3
    for kind, title in (("improve_product", "Что улучшить в товаре"), ("highlight_in_card", "Что подчеркнуть в карточке")):
        ws.cell(row=row, column=1, value=title).font = TITLE
        row += 1
        _header(ws, row, ["Совет", "Подробно", "На основе тем"])
        row += 1
        for r in summary["recommendations"][kind]:
            _row(ws, row, [r["title"], r["details"], ", ".join(r.get("based_on") or [])])
            row += 1
        row += 1
    ws.cell(row=row, column=1, value="Кратко").font = TITLE
    c = ws.cell(row=row + 1, column=1, value=summary["chat_summary"])
    c.alignment = WRAP
    ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=3)
    ws.row_dimensions[row + 1].height = 15 * (summary["chat_summary"].count("\n") + len(summary["chat_summary"]) // 120 + 1)


def write_competitor_sheet(ws: Worksheet, r: dict, m: dict) -> None:
    neg, pos = r["reviews_analyzed"]["negative"], r["reviews_analyzed"]["positive"]
    ws["A1"] = f"{r['brand']} — {r['name']}".strip(" —")
    ws["A1"].font = TITLE
    url = m.get("url") or f"https://www.wildberries.ru/catalog/{r['article']}/detail.aspx"
    ws["A2"] = url
    ws["A2"].hyperlink = url
    ws["A2"].style = "Hyperlink"
    info = [f"Артикул: {r['article']}"]
    if m.get("rating"):
        info.append(f"рейтинг {m['rating']}")
    if m.get("total_feedbacks"):
        info.append(f"всего оценок {m['total_feedbacks']}")
    info.append(f"проанализировано: {neg} негативных, {pos} позитивных")
    if m.get("months"):
        info.append(f"за {m['months']} мес.")
    ws["A3"] = ", ".join(info)
    _widths(ws, [34, 50, 13, 13, 50, 50, 50])
    row = 5
    for section, title, base in (("complaints", "Жалобы", neg), ("praise", "Похвала", pos)):
        ws.cell(row=row, column=1, value=title).font = TITLE
        row += 1
        _header(ws, row, ["Тема", "Описание", "Упоминаний", "Доля, %", "Пример 1", "Пример 2", "Пример 3"])
        row += 1
        if not r[section]:
            ws.cell(row=row, column=1, value="— нет —")
            row += 1
        for t in r[section]:
            share = round(100 * t["mentions"] / base) if base else None
            examples = [f"«{e['quote']}» ({e['rating']}★)" for e in t["examples"]]
            _row(ws, row, [t["theme"], t["description"], t["mentions"], share, *examples])
            row += 1
        row += 1


def write_status_sheet(ws: Worksheet, status: list[dict], results: dict) -> None:
    ws["A1"] = "Сбор и анализ по товарам"
    ws["A1"].font = TITLE
    _widths(ws, [16, 22, 18, 70])
    _header(ws, 3, ["Артикул", "Сбор", "Анализ", "Комментарий"])
    for i, s in enumerate(status, 4):
        art = s.get("article")
        analysed = "есть" if art in results else ("—" if s["status"] != "ok" else "нет (ошибка анализа)")
        comment = s.get("message") or ""
        if not art:
            comment = f"{s.get('input', '')}: {comment}"
        _row(ws, i, [art or "", STATUS_RU.get(s["status"], s["status"]), analysed, comment])


def build_report(run: Path, out: Path, log=print) -> dict:
    results, meta, summary, status = load_run(run, log)
    if not results:
        raise SystemExit("Нет ни одного валидного results/<артикул>.json — отчёт строить не из чего")
    wb = Workbook()
    write_summary_sheet(wb.active, results, meta, summary)
    used = {"свод", "рекомендации", "сбор"}
    write_recommendations_sheet(wb.create_sheet("Рекомендации"), summary)
    # листы конкурентов — в порядке ссылок из входного файла
    order = [s["article"] for s in status if s.get("article") in results]
    order += [a for a in results if a not in order]
    for art in dict.fromkeys(order):
        write_competitor_sheet(wb.create_sheet(sheet_title(_label(art, results, meta), used)), results[art], meta.get(art, {}))
    write_status_sheet(wb.create_sheet("Сбор"), status, results)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return {"results": results, "summary": summary, "status": status}


def console_summary(info: dict, out: Path) -> str:
    results, summary = info["results"], info["summary"]
    lines = [f"Отчёт: {out}", f"Конкурентов в отчёте: {len(results)}"]
    if summary:
        top = sorted(((schema.summary_stats(g, results, "complaints"), g["theme"]) for g in summary["complaints"]),
                     key=lambda x: (-x[0]["competitors"], -x[0]["mentions"]))[:5]
        lines.append("Топ жалоб (упоминаний / конкурентов):")
        lines += [f"  - {t}: {st['mentions']} / {st['competitors']}" + (" — общая для ниши" if st["niche_wide"] else "")
                  for st, t in top]
    else:
        lines.append("Сводки по нише нет — только листы конкурентов.")
    problems = [s for s in info["status"] if s["status"] != "ok" or s.get("article") not in results]
    if problems:
        lines.append("Не вошли в отчёт:")
        for s in problems:
            why = s.get("message") or ("ошибка анализа" if s["status"] == "ok" else s["status"])
            lines.append(f"  - {s.get('article') or s.get('input')}: {why}")
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Excel-отчёт по отзывам конкурентов")
    p.add_argument("--run", type=Path, default=Path("."), help="папка запуска с data/ и results/")
    p.add_argument("--out", type=Path, help="путь к xlsx (по умолчанию <run>/report.xlsx)")
    a = p.parse_args(argv)
    out = a.out or a.run / "report.xlsx"
    try:
        info = build_report(a.run, out)
    except SystemExit as e:
        print(f"Ошибка: {e}", file=sys.stderr)
        return 1
    print(console_summary(info, out))
    return 0


if __name__ == "__main__":
    sys.exit(main())
